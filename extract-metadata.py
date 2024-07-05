import os
import sys
from docx import Document
from openpyxl import load_workbook
import PyPDF2

def get_docx_metadata(file_path):
    try:
        doc = Document(file_path)
        core_properties = doc.core_properties
        metadata = {
            "Title": core_properties.title,
            "Author": core_properties.author,
            "Subject": core_properties.subject,
            "Keywords": core_properties.keywords,
            "Last Modified By": core_properties.last_modified_by,
            "Created": core_properties.created,
            "Modified": core_properties.modified
        }
        return metadata
    except Exception as e:
        print(f"Error al leer {file_path}: {e}")
        return {}

def get_xlsx_metadata(file_path):
    try:
        workbook = load_workbook(filename=file_path)
        core_properties = workbook.properties
        metadata = {
            "Title": core_properties.title,
            "Author": core_properties.creator,
            "Subject": core_properties.subject,
            "Keywords": core_properties.keywords,
            "Last Modified By": core_properties.last_modified_by,
            "Created": core_properties.created,
            "Modified": core_properties.modified
        }
        return metadata
    except Exception as e:
        print(f"Error al leer {file_path}: {e}")
        return {}

def get_pdf_metadata(file_path):
    try:
        with open(file_path, 'rb') as file:
            reader = PyPDF2.PdfFileReader(file)
            doc_info = reader.getDocumentInfo()
            metadata = {
                "Title": doc_info.title,
                "Author": doc_info.author,
                "Subject": doc_info.subject,
                "Producer": doc_info.producer,
            }
            # Añadir metadatos opcionales si están presentes
            optional_metadata = {
                "Keywords": doc_info.get('/Keywords'),
                "Creator": doc_info.get('/Creator'),
                "CreationDate": doc_info.get('/CreationDate'),
                "ModDate": doc_info.get('/ModDate'),
            }
            metadata.update({k: v for k, v in optional_metadata.items() if v})
        return metadata
    except Exception as e:
        print(f"Error al leer {file_path}: {e}")
        return {}

def extract_metadata(directory):
    for root, dirs, files in os.walk(directory):
        for file in files:
            file_path = os.path.join(root, file)
            print(f"Procesando {file_path}...")
            if file.endswith('.docx'):
                print(f"\nMetadatos de {file}:")
                metadata = get_docx_metadata(file_path)
                for key, value in metadata.items():
                    print(f"{key}: {value}")
            elif file.endswith('.xlsx'):
                print(f"\nMetadatos de {file}:")
                metadata = get_xlsx_metadata(file_path)
                for key, value in metadata.items():
                    print(f"{key}: {value}")
            elif file.endswith('.pdf'):
                print(f"\nMetadatos de {file}:")
                metadata = get_pdf_metadata(file_path)
                for key, value in metadata.items():
                    print(f"{key}: {value}")

if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Uso: python extract_metadata.py <ruta_del_directorio>")
        sys.exit(1)
    
    directory = sys.argv[1]
    if not os.path.isdir(directory):
        print(f"Error: {directory} no es un directorio válido")
        sys.exit(1)

    extract_metadata(directory)
